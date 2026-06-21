import os
import sys
import re
import json
import webbrowser
import http.server
import socketserver
import urllib.parse
import difflib
from pathlib import Path

# Try importing dependencies
try:
    import docx
except ImportError:
    print("Warning: python-docx is not installed. Word documents won't be parsed.")
    docx = None

try:
    import pypdf
except ImportError:
    print("Warning: pypdf is not installed. PDF documents won't be parsed.")
    pypdf = None

try:
    import fitz
except ImportError:
    print("Warning: pymupdf is not installed. PDF styled text parsing is disabled.")
    fitz = None

# Configure encoding
if sys.stdout:
    sys.stdout.reconfigure(encoding='utf-8')

# Try relative path first for portability, fallback to user's absolute desktop path
RELATIVE_ON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TaiLieuOnTap")
if os.path.exists(RELATIVE_ON):
    ON_FOLDER = RELATIVE_ON
else:
    ON_FOLDER = r"c:\Users\MSIIIIII\Desktop\On"
PORT = 8000

def clean_vietnamese_splits(t):
    p1 = r'\b(b|c|ch|d|đ|g|gh|gi|h|k|kh|l|m|n|ng|ngh|nh|p|ph|qu|r|s|t|th|tr|v|x)\s+([àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđĐ]+[a-zA-Z]*)\b'
    p2 = r'\b(chuy|nguy|tuy|khuy|truy|duy|huy|luy|quy|khi|ti|bi|di|li|ni|si|vi|ki|chi|hi|phi|thi|tri|xuy|mu)\s+(ên|ển|ết|ếp|êu|êm|ệm|ệt|ến|ền|ễn|ệu|êu|ốn|ôn|ồn|ộn|ổn|an|án|àn|ạn|ản|ãn)\b'
    def repl(m):
        return m.group(1) + m.group(2)
    old_t = ""
    while old_t != t:
        old_t = t
        t = re.sub(p1, repl, t, flags=re.IGNORECASE)
        t = re.sub(p2, repl, t, flags=re.IGNORECASE)
    return t

def format_numbered_lists(text):
    if re.search(r'\b1\.\s*\S', text) and re.search(r'\b2\.\s*\S', text):
        text = re.sub(r'\s+(\d+)\.\s*', r'\n\1. ', text)
        text = re.sub(r'^(.*?)\s+1\.\s*', r'\1\n1. ', text)
    return text

def clean_text(t, is_option=False):
    t = re.sub(r'\s+', ' ', t).strip()
    t = clean_vietnamese_splits(t)
    if not is_option:
        t = format_numbered_lists(t)
    return t

def clean_opt(o):
    o = re.sub(r"^[A-D]\.\s*", "", o, flags=re.IGNORECASE)
    o = re.sub(r"\s+", " ", o).strip().lower()
    return o

def get_clean_key(t):
    t = t.lower()
    t = re.sub(r"^câu\s+\d+[:.]?\s*", "", t)
    t = re.sub(r"^\d+[:.]?\s*", "", t)
    t = re.sub(r'[^a-z0-9àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ]', '', t)
    return t

def get_very_clean_key(t):
    t = t.lower()
    t = re.sub(r"^câu\s+\d+[:.]?\s*", "", t)
    t = re.sub(r"^[a-d]\.\s*", "", t)
    t = re.sub(r'[^a-z0-9àáạảãâầấậẩẫăằắặẳẵèéẹẻẽêềếệểễìíịỉĩòóọỏõôồốộổỗơờớợởỡùúụủũưừứựửữỳýỵỷỹđ]', '', t)
    return t

# Parser for Trac nghiem QTHTM.docx
def parse_trac_nghiem(path):
    if not docx:
        return []
    doc = docx.Document(path)
    questions = []
    curr = None
    for p in doc.paragraphs:
        p_lines = p.text.split('\n')
        for t in p_lines:
            t = t.strip()
            if not t:
                continue
            if re.match(r"^Câu\s+\d+", t, re.IGNORECASE):
                if curr:
                    questions.append(curr)
                curr = {"question": t, "options": [], "answer": "", "source": os.path.basename(path)}
            elif curr:
                opts = re.findall(r"((?:^|\s+)[A-D]\.\s+.*?)(?=\s+[A-D]\.\s+|\t[A-D]\.\s+|$)", t)
                if opts:
                    curr["options"].extend([o.strip() for o in opts])
                elif re.match(r"^[A-D]\.", t):
                    curr["options"].append(t)
                elif "Đáp án:" in t or "Dap an:" in t:
                    m = re.search(r"(?:Đáp án|Dap an):\s*([A-D])", t, re.IGNORECASE)
                    if m:
                        curr["answer"] = m.group(1).upper()
                else:
                    if curr["options"]:
                        curr["options"][-1] += " " + t
                    else:
                        curr["question"] += " " + t
    if curr:
        questions.append(curr)
    
    for q in questions:
        q["question"] = clean_text(q["question"])
        q["options"] = [clean_text(o, is_option=True) for o in q["options"]]
    return questions

# Parser for QTHTM.docx
def parse_qthtm(path):
    if not docx:
        return []
    doc = docx.Document(path)
    questions = []
    curr = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if re.match(r"^Câu\s+\d+[:.]", t, re.IGNORECASE):
            if curr:
                questions.append(curr)
            curr = {"question": t, "options": [], "answer": "", "source": os.path.basename(path)}
        elif curr:
            opts = re.findall(r"((?:^|\s+)[A-D]\.\s+.*?)(?=\s+[A-D]\.\s+|\t[A-D]\.\s+|$)", t)
            if opts:
                for opt in opts:
                    opt = opt.strip()
                    curr["options"].append(opt)
                    
                    # Prioritize highlight over bold
                    has_highlight = any(r.font.highlight_color is not None for r in p.runs)
                    marked_text = ""
                    for r in p.runs:
                        if r.text:
                            if has_highlight:
                                if r.font.highlight_color is not None:
                                    marked_text += r.text
                            else:
                                if r.bold:
                                    marked_text += r.text
                                    
                    if marked_text:
                        m_ans = re.match(r"^([A-D])\.", opt)
                        if m_ans:
                            opt_clean = get_very_clean_key(opt)
                            marked_clean = get_very_clean_key(marked_text)
                            if opt_clean in marked_clean or (marked_clean and marked_clean in opt_clean):
                                curr["answer"] = m_ans.group(1).upper()
            else:
                m = re.match(r"^([A-D])\.", t)
                if m:
                    curr["options"].append(t)
                    marked_len = 0
                    total_len = 0
                    has_highlight = any(r.font.highlight_color is not None for r in p.runs)
                    for r in p.runs:
                        rt = r.text.strip()
                        if rt:
                            total_len += len(rt)
                            if has_highlight:
                                if r.font.highlight_color is not None:
                                    marked_len += len(rt)
                            else:
                                if r.bold:
                                    marked_len += len(rt)
                    if total_len > 0 and (marked_len / total_len) > 0.5:
                        curr["answer"] = m.group(1).upper()
                else:
                    if curr["options"]:
                        curr["options"][-1] += " " + t
                    else:
                        curr["question"] += " " + t
    if curr:
        questions.append(curr)
        
    for q in questions:
        q["question"] = clean_text(q["question"])
        q["options"] = [clean_text(o, is_option=True) for o in q["options"]]
    return questions

# General PDF Parser for all three PDFs
def parse_pdf(path):
    if not pypdf:
        return []
    reader = pypdf.PdfReader(path)
    questions = []
    curr_q = None
    
    full_text = ""
    for page in reader.pages:
        full_text += (page.extract_text() or "") + "\n"
        
    count_cau = len(re.findall(r"(?:^|\n)Câu\s+\d+", full_text, re.IGNORECASE))
    count_num = len(re.findall(r"(?:^|\n)\d+\.\s+", full_text))
    format_type = 'num' if count_num > count_cau else 'cau'
    
    next_expected = 1
    for page_idx, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        lines = page_text.splitlines()
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            m = None
            if format_type == 'cau':
                m = re.match(r"^Câu\s+(\d+)[:.]?\s*(.*)", line, re.IGNORECASE)
            else:
                m = re.match(r"^(\d+)\.\s+(.*)", line)
                
            is_new_q = False
            if m:
                q_num = int(m.group(1))
                if format_type == 'num':
                    if q_num == next_expected:
                        is_new_q = True
                else:
                    is_new_q = True
                    
            if is_new_q:
                if curr_q:
                    questions.append(curr_q)
                curr_q = {
                    "num": int(m.group(1)),
                    "question": m.group(2),
                    "options": [],
                    "answer": "",
                    "source": os.path.basename(path),
                    "page_idx": page_idx,
                    "label": f"Câu {m.group(1)}" if format_type == 'cau' else f"{m.group(1)}."
                }
                if format_type == 'num':
                    next_expected = int(m.group(1)) + 1
            elif curr_q:
                opts = re.findall(r"((?:^|\s+)[A-D]\.\s+.*?)(?=\s+[A-D]\.\s+|\t[A-D]\.\s+|$)", line)
                if opts:
                    curr_q["options"].extend([o.strip() for o in opts])
                else:
                    if curr_q["options"]:
                        curr_q["options"][-1] += " " + line
                    else:
                        curr_q["question"] += " " + line
                        
    if curr_q:
        questions.append(curr_q)
        
    for q in questions:
        q["question"] = clean_text(q["question"])
        q["options"] = [clean_text(o, is_option=True) for o in q["options"]]
        
    return questions

# Build questions JSON database
def build_database():
    print(f"Scanning folder: {ON_FOLDER}...")
    if not os.path.exists(ON_FOLDER):
        print(f"Error: Folder {ON_FOLDER} does not exist.")
        return []

    docx1 = []
    docx2 = []
    pdf_on = []
    pdf_sai = []
    pdf_nd = []

    # Parse all files
    for f in os.listdir(ON_FOLDER):
        path = os.path.join(ON_FOLDER, f)
        if f == "Trac nghiem QTHTM.docx":
            docx1 = parse_trac_nghiem(path)
        elif f == "QTHTM.docx":
            docx2 = parse_qthtm(path)
        elif f == "ÔN TẬP CHO SV.pdf":
            pdf_on = parse_pdf(path)
        elif f == "SAI LA TOI - QTHTM.pdf":
            pdf_sai = parse_pdf(path)
        elif f == "NOI DUNG ON TAP - QTHTM - 04-6-2026.pdf":
            pdf_nd = parse_pdf(path)

    print(f"Found counts: Trac nghiem={len(docx1)}, QTHTM={len(docx2)}, ON TAP SV={len(pdf_on)}, SAI LA TOI={len(pdf_sai)}, NOI DUNG ON TAP={len(pdf_nd)}")

    # Build database of answers from docx
    db_answers = {}
    for q in docx1 + docx2:
        q_body = re.sub(r"^Câu\s+\d+[:.]?\s*", "", q["question"], flags=re.IGNORECASE)
        clean_k = get_clean_key(q_body)
        if q["answer"]:
            db_answers[clean_k] = {
                "answer": q["answer"],
                "options": q["options"]
            }

    # Manual answers for NOI DUNG ON TAP - QTHTM - 04-6-2026.pdf
    manual_answers_nd = {
        2: "C", 3: "C", 5: "B", 16: "D", 34: "D", 38: "C", 46: "D", 62: "C", 63: "C", 
        67: "D", 69: "D", 76: "D", 77: "C", 80: "D", 82: "D", 83: "B", 84: "B", 
        85: "B", 87: "B", 102: "B", 103: "D", 104: "D", 110: "B", 127: "B", 
        132: "C", 133: "C", 134: "D", 135: "D", 151: "B", 155: "C", 173: "D", 
        174: "B", 175: "B"
    }

    # Extract red-colored answers directly from PDFs using PyMuPDF (fitz)
    # Extract red-colored answers directly from PDFs using PyMuPDF (fitz)
    def extract_red_answers_from_pdf(path, questions):
        if not fitz:
            return
        try:
            def clean_span(s):
                return re.sub(r"\s+", " ", s).strip().lower()

            doc = fitz.open(path)
            
            # Find coordinates of each question in the PDF by searching for label (unique on each page)
            q_coords = []
            page_label_counters = {} # key: (page_idx, label), value: occurrence count
            
            for q in questions:
                page_idx = q.get("page_idx", 0)
                page = doc[page_idx]
                
                label = q.get("label", f"Câu {q['num']}")
                rects = page.search_for(label)
                
                valid_rects = [r for r in rects if r.x0 < 100]
                if not valid_rects:
                    valid_rects = rects
                    
                key = (page_idx, label)
                occurrence = page_label_counters.get(key, 0)
                
                best_rect = None
                if valid_rects:
                    if occurrence < len(valid_rects):
                        best_rect = valid_rects[occurrence]
                    else:
                        best_rect = valid_rects[-1]
                        
                page_label_counters[key] = occurrence + 1
                
                if best_rect:
                    q_coords.append({
                        "q": q,
                        "page": page_idx,
                        "y": best_rect.y0
                    })
                else:
                    # Fallback to search using question text prefix
                    q_prefix = q["question"][:30].strip()
                    rects = page.search_for(q_prefix)
                    if rects:
                        q_coords.append({
                            "q": q,
                            "page": page_idx,
                            "y": rects[0].y0
                        })
                    else:
                        q_coords.append({
                            "q": q,
                            "page": page_idx,
                            "y": 0.0
                        })
                        
            def extract_red_spans_in_range(page, y_start, y_end):
                spans = []
                blocks = page.get_text("dict")["blocks"]
                for b in blocks:
                    if "lines" not in b:
                        continue
                    for l in b["lines"]:
                        line_y = l["bbox"][1]
                        if y_start <= line_y < y_end:
                            for s in l["spans"]:
                                text = s["text"].strip()
                                color = s["color"]
                                if text:
                                    r = (color >> 16) & 255
                                    g = (color >> 8) & 255
                                    b_val = color & 255
                                    if r > 200 and g < 50 and b_val < 50:
                                        cleaned = clean_span(text)
                                        if cleaned:
                                            spans.append(cleaned)
                return spans

            def get_red_spans_for_question(item, q_coords, doc):
                q = item["q"]
                page_idx = item["page"]
                y_start = item["y"]
                
                red_spans = []
                
                next_q_on_same_page = None
                next_q_idx = None
                for i, other in enumerate(q_coords):
                    if other["q"] == q:
                        next_q_idx = i
                        break
                        
                for other in q_coords[next_q_idx+1:]:
                    if other["page"] == page_idx:
                        next_q_on_same_page = other
                        break
                        
                if next_q_on_same_page:
                    y_end = next_q_on_same_page["y"]
                    red_spans.extend(extract_red_spans_in_range(doc[page_idx], y_start, y_end))
                else:
                    # Question overflows page N to page N+1
                    y_end = doc[page_idx].rect.height
                    red_spans.extend(extract_red_spans_in_range(doc[page_idx], y_start, y_end))
                    
                    if page_idx + 1 < len(doc):
                        next_page_idx = page_idx + 1
                        first_q_on_next_page = None
                        for other in q_coords:
                            if other["page"] == next_page_idx:
                                first_q_on_next_page = other
                                break
                        
                        if first_q_on_next_page:
                            y_end_next = first_q_on_next_page["y"]
                        else:
                            y_end_next = doc[next_page_idx].rect.height
                            
                        red_spans.extend(extract_red_spans_in_range(doc[next_page_idx], 0, y_end_next))
                        
                return red_spans

            # Resolve answers for each question
            for item in q_coords:
                q = item["q"]
                
                red_spans = get_red_spans_for_question(item, q_coords, doc)
                
                # Pass 1: Try prefix matching first (highly accurate, avoids leakage and substring issues)
                matched_letter = None
                for opt in q["options"]:
                    m_ans = re.match(r"^([A-D])\.", opt)
                    if not m_ans:
                        continue
                    letter = m_ans.group(1).upper()
                    letter_lower = letter.lower()
                    opt_clean_full = clean_span(opt)
                    
                    for red in red_spans:
                        if red == f"{letter_lower}." or red == letter_lower:
                            matched_letter = letter
                            break
                        if red.startswith(f"{letter_lower}.") or red.startswith(f"{letter_lower} "):
                            if opt_clean_full.startswith(red) or red.startswith(opt_clean_full):
                                matched_letter = letter
                                break
                    if matched_letter:
                        break
                        
                if matched_letter:
                    q["answer"] = matched_letter
                else:
                    # Pass 2: Fallback to substring matching if no prefix matched
                    for opt in q["options"]:
                        m_ans = re.match(r"^([A-D])\.", opt)
                        if not m_ans:
                            continue
                        letter = m_ans.group(1).upper()
                        opt_clean_val = clean_opt(opt)
                        
                        matched = False
                        for red in red_spans:
                            if len(red) >= 10 and (red == opt_clean_val or red in opt_clean_val or opt_clean_val in red):
                                matched = True
                                break
                        if matched:
                            q["answer"] = letter
                            break
        except Exception as e:
            print(f"Warning: Failed to extract red answers from {os.path.basename(path)}: {e}")

    # Extract red answers first (highly accurate direct extraction)
    extract_red_answers_from_pdf(os.path.join(ON_FOLDER, "SAI LA TOI - QTHTM.pdf"), pdf_sai)
    extract_red_answers_from_pdf(os.path.join(ON_FOLDER, "NOI DUNG ON TAP - QTHTM - 04-6-2026.pdf"), pdf_nd)
    extract_red_answers_from_pdf(os.path.join(ON_FOLDER, "ÔN TẬP CHO SV.pdf"), pdf_on)

    # Helper function to match answers from the database
    def resolve_answers(qs, is_nd=False):
        for q in qs:
            # Skip if answer was already extracted from red text
            if q["answer"]:
                continue
            if is_nd and q["num"] in manual_answers_nd:
                q["answer"] = manual_answers_nd[q["num"]]
                continue
                
            q_body = re.sub(r"^Câu\s+\d+[:.]?\s*", "", q["question"], flags=re.IGNORECASE)
            clean_k = get_clean_key(q_body)
            
            # Try exact match first, then substring match
            matched_db_q = None
            if clean_k in db_answers:
                matched_db_q = db_answers[clean_k]
            else:
                for dk in db_answers:
                    if dk in clean_k or clean_k in dk:
                        matched_db_q = db_answers[dk]
                        break
            
            if matched_db_q:
                q["answer"] = matched_db_q["answer"]
                if not is_nd and (not q["options"] or len(q["options"]) < 4 or any(len(o) < 5 for o in q["options"])):
                    q["options"] = matched_db_q["options"]
            else:
                # Try option set matching
                pq_opts = sorted([clean_opt(o) for o in q["options"]])
                if not pq_opts:
                    continue
                best_match = None
                best_score = 0
                for dq in docx1 + docx2:
                    dq_opts = sorted([clean_opt(o) for o in dq["options"]])
                    if not dq_opts:
                        continue
                    score = 0
                    for po in pq_opts:
                        if any(difflib.get_close_matches(po, [do], n=1, cutoff=0.7) for do in dq_opts):
                            score += 1
                    if score > best_score:
                        best_score = score
                        best_match = dq
                
                if best_match and best_score >= min(len(pq_opts), 3):
                    q["answer"] = best_match["answer"]
                    if not is_nd:
                        q["options"] = best_match["options"]

    # Match remaining answers for the three PDFs
    resolve_answers(pdf_on)
    resolve_answers(pdf_sai)
    resolve_answers(pdf_nd, is_nd=True)

    # Combine into unified JSON database
    all_questions = []

    for idx, q in enumerate(docx1):
        all_questions.append({
            "id": f"docx1_{idx}",
            "num": idx + 1,
            "question": q["question"],
            "options": q["options"],
            "answer": q["answer"],
            "source": q["source"]
        })

    for idx, q in enumerate(docx2):
        all_questions.append({
            "id": f"docx2_{idx}",
            "num": idx + 1,
            "question": q["question"],
            "options": q["options"],
            "answer": q["answer"],
            "source": q["source"]
        })

    for idx, q in enumerate(pdf_on):
        if q["answer"]:
            all_questions.append({
                "id": f"pdf_on_{idx}",
                "num": q["num"],
                "question": q["question"],
                "options": q["options"],
                "answer": q["answer"],
                "source": q["source"]
            })

    for idx, q in enumerate(pdf_sai):
        if q["answer"]:
            all_questions.append({
                "id": f"pdf_sai_{idx}",
                "num": q["num"],
                "question": q["question"],
                "options": q["options"],
                "answer": q["answer"],
                "source": q["source"]
            })

    for idx, q in enumerate(pdf_nd):
        if q["answer"]:
            all_questions.append({
                "id": f"pdf_nd_{idx}",
                "num": q["num"],
                "question": q["question"],
                "options": q["options"],
                "answer": q["answer"],
                "source": q["source"]
            })

    # Save to files (JSON and JS fallback for file:// protocol)
    web_dir = Path(__file__).parent / "web"
    web_dir.mkdir(exist_ok=True)
    
    json_path = web_dir / "questions.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(all_questions, f, ensure_ascii=False, indent=2)
        
    js_path = web_dir / "questions.js"
    with open(js_path, "w", encoding="utf-8") as f:
        f.write("window.allQuestionsData = ")
        json.dump(all_questions, f, ensure_ascii=False, indent=2)
        f.write(";")
        
    print(f"Generated {len(all_questions)} questions in {json_path} and {js_path}")
    return all_questions

# HTTP Request Handler
class QuizRequestHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(Path(__file__).parent / "web"), **kwargs)
        
    def end_headers(self):
        self.send_header('Cache-Control', 'no-store, no-cache, must-revalidate, max-age=0')
        self.send_header('Pragma', 'no-cache')
        self.send_header('Expires', '0')
        super().end_headers()

def run_server():
    # Build database on startup
    build_database()

    # Start Server
    handler = QuizRequestHandler
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("", PORT), handler) as httpd:
        print(f"Server starting on http://localhost:{PORT}")
        webbrowser.open(f"http://localhost:{PORT}")
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nServer stopped.")

if __name__ == "__main__":
    run_server()
